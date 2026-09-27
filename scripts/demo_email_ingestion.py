"""Deterministic email demo. Creates synthetic cases via n8n or the direct API."""

import argparse
import base64
import json
import uuid
from io import BytesIO
from pathlib import Path

import httpx
from openpyxl import load_workbook


def fixture(run_id: str, variant: str, attachments: list[dict]) -> dict:
    return {
        "external_message_id": f"<boah-demo-{run_id}-{variant}@example.test>",
        "sender": {"address": "jan@example.test", "name": "Jan Kowalski"},
        "recipients": [{"address": "office@example.test", "name": "Example office"}],
        "subject": "Zapytanie ofertowe — modernizacja biura",
        "received_at": "2026-09-27T12:00:00+02:00",
        "sent_at": "2026-09-27T09:59:00Z",
        "text_body": "Dzień dobry,\nproszę o przygotowanie wyceny zgodnie z załączoną specyfikacją.",
        "html_body": "<p>Synthetic source HTML — displayed as escaped text.</p>",
        "attachments": attachments,
    }


def as_n8n(payload: dict) -> dict:
    return {
        "mail": {
            "messageId": payload["external_message_id"],
            "from": {"value": [payload["sender"]]},
            "to": {"value": payload["recipients"]},
            "subject": payload["subject"],
            "receivedAt": payload["received_at"],
            "date": payload["sent_at"],
            "text": payload["text_body"],
            "html": payload["html_body"],
        },
        "attachments": payload["attachments"],
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--via", choices=["n8n", "api"], default="n8n")
    parser.add_argument("--run-id", default=uuid.uuid4().hex[:10])
    parser.add_argument("--api-url", default="http://localhost:8000/api/v1")
    parser.add_argument(
        "--webhook-url", default="http://localhost:5678/webhook/boah-email-fixture-m2"
    )
    parser.add_argument("--output-dir", type=Path)
    args = parser.parse_args()
    sample = (
        Path(__file__).resolve().parents[1] / "sample_data" / "email_specification.txt"
    )
    attachment = {
        "filename": sample.name,
        "mime_type": "text/plain",
        "content_base64": base64.b64encode(sample.read_bytes()).decode(),
    }
    extra = {
        "filename": "extra.txt",
        "mime_type": "text/plain",
        "content_base64": base64.b64encode(
            b"Description: Synthetic additional document"
        ).decode(),
    }
    variants = {
        "zero": [],
        "one": [attachment],
        "multiple": [attachment, {**attachment, "filename": "renamed.txt"}, extra],
    }
    cases = []
    with httpx.Client(timeout=65) as client:

        def deliver(payload):
            response = client.post(
                args.webhook_url
                if args.via == "n8n"
                else args.api_url + "/inbound/email",
                json=as_n8n(payload) if args.via == "n8n" else payload,
            )
            response.raise_for_status()
            return response.json()

        before = client.get(args.api_url + "/cases").json()["total"]
        for variant, files in variants.items():
            payload = fixture(args.run_id, variant, files)
            first = deliver(payload)
            assert first["result"] == "created", (
                "Use a new --run-id for a new measured run"
            )
            replay = deliver(payload)
            assert (
                replay["result"] == "duplicate"
                and replay["case_id"] == first["case_id"]
            )
            case = client.get(args.api_url + "/cases/" + first["case_id"]).json()
            assert (
                case["source"] == "email"
                and case["inbound_message"]["text_body"] == payload["text_body"]
            )
            assert (
                len(case["attachments"])
                == {"zero": 0, "one": 1, "multiple": 2}[variant]
            )
            assert (
                case["status"] == "REVIEW_REQUIRED"
            )  # M1 treats .test as a reserved email domain.
            assert any(
                e["event_type"] == "EMAIL_DUPLICATE_IGNORED"
                for e in case["audit_events"]
            )
            cases.append(
                {
                    "variant": variant,
                    "case_id": case["id"],
                    "public_case_id": case["public_id"],
                    "message_id": first["message_id"],
                    "attachments": len(case["attachments"]),
                    "status": case["status"],
                    "audit_events": len(case["audit_events"]),
                }
            )
        after = client.get(args.api_url + "/cases").json()["total"]
        assert after - before == 3
        chosen = cases[1]
        identifier = chosen["case_id"]
        assert (
            client.post(f"{args.api_url}/review/{identifier}/approve").status_code
            == 409
        )
        corrected = client.patch(
            f"{args.api_url}/review/{identifier}",
            json={"customer_email": "jan@example.com"},
        )
        corrected.raise_for_status()
        assert corrected.json()["status"] == "READY"
        approved = client.post(f"{args.api_url}/review/{identifier}/approve")
        approved.raise_for_status()
        assert approved.json()["status"] == "APPROVED"
        export_ids = {}
        for kind in ("json", "xlsx"):
            response = client.post(f"{args.api_url}/exports/{identifier}/{kind}")
            response.raise_for_status()
            export_ids[kind] = response.headers["x-export-id"]
            if kind == "json":
                assert response.json()["data"]["estimated_value"] == "12500.00"
            else:
                book = load_workbook(BytesIO(response.content))
                assert book.sheetnames == [
                    "Summary",
                    "Attachments",
                    "Validation",
                    "Audit",
                ]
                assert dict(book["Summary"].values)["Estimated Value"] == 12500
            if args.output_dir:
                args.output_dir.mkdir(parents=True, exist_ok=True)
                (args.output_dir / f"{chosen['public_case_id']}.{kind}").write_bytes(
                    response.content
                )
        final_replay = deliver(fixture(args.run_id, "one", variants["one"]))
        assert (
            final_replay["result"] == "duplicate"
            and final_replay["status"] == "EXPORTED"
        )
        final = client.get(f"{args.api_url}/cases/{identifier}").json()
        chosen.update(
            status=final["status"],
            audit_events=len(final["audit_events"]),
            exports=export_ids,
        )
        report = {
            "result": "PASS",
            "via": args.via,
            "run_id": args.run_id,
            "deliveries": 7,
            "new_cases": after - before,
            "additional_cases_from_replay": 0,
            "stored_attachments": sum(c["attachments"] for c in cases),
            "cases": cases,
        }
        if args.output_dir:
            (args.output_dir / "milestone-2-email-demo.json").write_text(
                json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8"
            )
        print(json.dumps(report, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
