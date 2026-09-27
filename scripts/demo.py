"""Repeatable HTTP E2E demo against the running Compose application (creates demo data)."""

import argparse
import hashlib
import json
from io import BytesIO
from pathlib import Path

import httpx
from openpyxl import load_workbook


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base-url", default="http://localhost:8000/api/v1")
    parser.add_argument("--output-dir", type=Path)
    args = parser.parse_args()
    fixture = (
        Path(__file__).resolve().parents[1] / "sample_data" / "review_required.txt"
    )
    content = fixture.read_bytes()
    with httpx.Client(base_url=args.base_url, timeout=30) as client:
        response = client.post("/cases", json={"source": "api"})
        response.raise_for_status()
        case = response.json()
        identifier = case["id"]
        response = client.post(
            f"/uploads/{identifier}",
            files={"files": (fixture.name, content, "text/plain")},
        )
        response.raise_for_status()
        case = response.json()
        assert case["status"] == "REVIEW_REQUIRED"
        assert client.post(f"/review/{identifier}/approve").status_code == 409
        assert client.post(f"/exports/{identifier}/json").status_code == 409
        duplicate = client.post(
            f"/uploads/{identifier}",
            files={"files": ("renamed.txt", content, "text/plain")},
        )
        duplicate.raise_for_status()
        assert len(duplicate.json()["attachments"]) == 1
        raw_fields = duplicate.json()["extracted_fields"]
        corrected = client.patch(
            f"/review/{identifier}",
            json={"tax_id": "5260250274", "estimated_value": "12 500,00 PLN"},
        )
        corrected.raise_for_status()
        assert corrected.json()["status"] == "READY"
        assert corrected.json()["extracted_fields"] == raw_fields
        approved = client.post(f"/review/{identifier}/approve")
        approved.raise_for_status()
        assert approved.json()["status"] == "APPROVED"
        artifacts = {}
        for kind in ("json", "xlsx"):
            exported = client.post(f"/exports/{identifier}/{kind}")
            exported.raise_for_status()
            export_id = exported.headers["x-export-id"]
            fetched = client.get(f"/exports/{export_id}")
            assert fetched.content == exported.content
            artifacts[kind] = {
                "id": export_id,
                "sha256": hashlib.sha256(exported.content).hexdigest(),
            }
            if kind == "json":
                assert exported.json()["data"]["estimated_value"] == "12500.00"
            else:
                workbook = load_workbook(BytesIO(exported.content))
                assert workbook.sheetnames == [
                    "Summary",
                    "Attachments",
                    "Validation",
                    "Audit",
                ]
                assert dict(workbook["Summary"].values)["Estimated Value"] == 12500
            if args.output_dir:
                args.output_dir.mkdir(parents=True, exist_ok=True)
                (args.output_dir / f"{case['public_id']}.{kind}").write_bytes(
                    exported.content
                )
        final = client.get(f"/cases/{identifier}").json()
        assert final["status"] == "EXPORTED"
        events = {a["event_type"] for a in final["audit_events"]}
        assert {
            "case_created",
            "duplicate_detected",
            "field_manually_corrected",
            "validation_issue_resolved",
            "case_approved",
            "export_generated",
        } <= events
        assert (
            client.patch(
                f"/review/{identifier}", json={"request_title": "tamper"}
            ).status_code
            == 409
        )
        attachment = final["attachments"][0]
        assert (
            client.get(f"/uploads/{identifier}/{attachment['id']}").content == content
        )
        report = {
            "result": "PASS",
            "case_id": identifier,
            "public_id": case["public_id"],
            "status": final["status"],
            "attachment_id": attachment["id"],
            "original_sha256": attachment["sha256"],
            "audit_events": len(final["audit_events"]),
            "exports": artifacts,
        }
        if args.output_dir:
            (args.output_dir / "demo-result.json").write_text(
                json.dumps(report, indent=2), encoding="utf-8"
            )
        print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
