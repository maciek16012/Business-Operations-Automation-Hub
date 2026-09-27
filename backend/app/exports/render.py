import json
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.serialization import value_json
from app.validation.normalization import FIELDS


def approved_payload(detail: dict) -> dict:
    return {
        "schema_version": 1,
        "case_id": detail["public_id"],
        "status": detail["status"],
        "data": {field: detail[field] for field in FIELDS},
    }


def json_export(detail: dict) -> bytes:
    return json.dumps(approved_payload(detail), ensure_ascii=False, indent=2).encode("utf-8")


def xlsx_export(detail: dict) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    assert summary is not None
    summary.title = "Summary"
    summary.append(["Field", "Value"])
    summary.append(["Case ID", detail["public_id"]])
    for field in FIELDS:
        value = detail[field]
        if field == "estimated_value" and value is not None:
            from decimal import Decimal

            value = Decimal(value)
        if field == "requested_deadline" and value:
            from datetime import date

            value = date.fromisoformat(value)
        summary.append([field.replace("_", " ").title(), value])
        if field == "estimated_value":
            summary.cell(summary.max_row, 2).number_format = "#,##0.00"
        if field == "requested_deadline":
            summary.cell(summary.max_row, 2).number_format = "yyyy-mm-dd"
    summary.append(["Status", detail["status"]])
    for title, key, columns in (
        (
            "Attachments",
            "attachments",
            ["original_filename", "mime_type", "size_bytes", "sha256", "created_at"],
        ),
        (
            "Validation",
            "validation_issues",
            ["code", "severity", "field_name", "message", "resolved", "created_at"],
        ),
        ("Audit", "audit_events", ["created_at", "event_type", "actor_type", "details"]),
    ):
        sheet = workbook.create_sheet(title)
        sheet.append([c.replace("_", " ").title() for c in columns])
        for row in detail[key]:
            sheet.append(
                [
                    json.dumps(row[c], ensure_ascii=False)
                    if isinstance(row[c], dict)
                    else value_json(row[c])
                    for c in columns
                ]
            )
    for sheet in workbook:
        for cell in sheet[1]:
            if cell.value == "Created At":
                for row in sheet.iter_rows(min_row=2, min_col=cell.column, max_col=cell.column):
                    timestamp = row[0]
                    if timestamp.value:
                        timestamp.value = datetime.fromisoformat(str(timestamp.value)).replace(
                            tzinfo=None
                        )
                        timestamp.number_format = "yyyy-mm-dd hh:mm:ss"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="243B53")
        for index, column in enumerate(sheet.columns, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = min(
                70, max(18, max(len(str(cell.value or "")) for cell in column) + 2)
            )
        for row in sheet.iter_rows():
            for cell in row:
                # Untrusted document text must never become spreadsheet formulas.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    content = BytesIO()
    workbook.save(content)
    return content.getvalue()
