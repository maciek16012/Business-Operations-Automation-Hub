import runpy
from pathlib import Path
from unittest.mock import AsyncMock

import pytest

from app.adaptive.classifier import LocalClassifier
from app.adaptive.contracts import Evidence
from app.adaptive.layout import LayoutDetector
from app.adaptive.validation import validate_table
from app.core.config import settings
from app.security.preflight import ClamAV

ROOT = Path(__file__).resolve().parents[2]
fixture = runpy.run_path(str(ROOT / "scripts/milestone6/fixtures.py"))["fixture"]


@pytest.fixture
def adaptive(monkeypatch):
    monkeypatch.setattr(settings, "adaptive_extraction_enabled", True)
    monkeypatch.setattr(settings, "security_preflight_enabled", True)
    monkeypatch.setattr(settings, "ocr_enabled", True)
    monkeypatch.setattr(settings, "stp_enabled", True)
    monkeypatch.setattr(ClamAV, "scan", AsyncMock(return_value=("ClamAV test-double", None)))


async def upload(client, kind, variant=0):
    response = await client.post(
        "/api/v1/cases",
        json={
            "customer_name": "Synthetic operator",
            "request_title": f"M6 {kind}",
            "customer_email": "m6@example.com",
        },
    )
    cid = response.json()["id"]
    content, truth = fixture(kind, variant)
    response = await client.post(
        f"/api/v1/uploads/{cid}", files={"files": ("document.pdf", content, "application/pdf")}
    )
    assert response.status_code == 200, response.text
    return response.json(), truth


@pytest.mark.parametrize(
    "kind", ["INVOICE", "PRINTED_TABLE", "HANDWRITTEN_TABLE", "GENERIC_DOCUMENT", "UNKNOWN"]
)
async def test_each_type_and_persistence(client, adaptive, kind):
    case, truth = await upload(client, kind)
    doc = case["documents"][0]
    assert doc["document_type"] == kind
    assert doc["classifier"] and doc["reasons"] and doc["created_at"]
    if kind != "INVOICE":
        assert case["tax_id"] is None and case["estimated_value"] is None
        assert case["status"] == "REVIEW_REQUIRED"
        assert case["review_tasks"][0]["task_type"] == "DOCUMENT_REVIEW"
        assert (await client.post(f"/api/v1/review/{case['id']}/approve")).status_code == 409
    else:
        assert case["tax_id"] == "5260250274"
        assert case["ocr_documents"] and doc["strategy"] == "existing-m4-invoice"
    if truth["tables"]:
        table = doc["tables"][0]
        assert table["row_count"] == truth["tables"][0]["row_count"]
        assert table["column_count"] == 6
        assert all(c["bbox"] for c in table["cells"])
        if kind == "HANDWRITTEN_TABLE":
            assert all(
                c["uncertain"] and c["raw_value"] is None for c in table["cells"] if c["row"] > 0
            )
    events = [e["event_type"] for e in case["audit_events"]]
    assert events.index("ATTACHMENT_SECURITY_SAFE") < events.index(
        "DOCUMENT_CLASSIFICATION_STARTED"
    )
    assert "ADAPTIVE_EXTRACTION_COMPLETED" in events


async def test_ambiguous_and_sparse_conservative():
    classifier = LocalClassifier()
    assert (
        classifier.classify(Evidence(text="invoice NIP VAT gross total assessment")).document_type
        == "UNKNOWN"
    )
    assert classifier.classify(Evidence(text="invoice")).review_required
    assert classifier.classify(Evidence(errors=["OCR unavailable"])).document_type == "UNKNOWN"


async def test_quarantine_never_classifies(client, adaptive, monkeypatch):
    sentinel = AsyncMock(side_effect=AssertionError("Unsafe document reached classifier"))
    monkeypatch.setattr(LayoutDetector, "detect", sentinel)
    cid = (await client.post("/api/v1/cases", json={})).json()["id"]
    result = await client.post(
        f"/api/v1/uploads/{cid}",
        files={"files": ("evil.pdf", b"MZ harmless sentinel", "application/pdf")},
    )
    assert result.status_code == 200
    assert not result.json()["documents"]
    sentinel.assert_not_called()


async def test_table_correction_revision_audit_and_approval(client, adaptive):
    case, truth = await upload(client, "HANDWRITTEN_TABLE")
    doc = case["documents"][0]
    path = f"/api/v1/documents/{doc['id']}/review"
    reason = {"reason": "Compared every cell against the synthetic original", "revision": 0}
    assert (await client.post(path, json=reason)).status_code == 422
    values = truth["tables"][0]["rows"]
    cells = [
        {"cell_id": c["id"], "value": values[c["row"]][c["column"]]}
        for c in doc["tables"][0]["cells"]
    ]
    bad = [dict(c) for c in cells]
    target = next(
        i for i, c in enumerate(doc["tables"][0]["cells"]) if c["row"] == 1 and c["column"] == 3
    )
    bad[target]["value"] = "999"
    assert (await client.post(path, json={**reason, "cells": bad})).status_code == 422
    reviewed = await client.post(path, json={**reason, "cells": cells})
    assert reviewed.status_code == 200, reviewed.text
    current = reviewed.json()
    assert current["documents"][0]["revision"] == 1
    assert current["review_tasks"][0]["status"] == "RESOLVED"
    assert current["status"] == "READY"
    assert any(e["event_type"] == "DOCUMENT_REVIEW_CORRECTED" for e in current["audit_events"])
    assert all(
        c["raw_value"] is None
        for c in current["documents"][0]["tables"][0]["cells"]
        if c["row"] > 0
    )
    assert (await client.post(path, json={**reason, "cells": cells})).status_code == 409
    assert (await client.post(f"/api/v1/review/{case['id']}/approve")).status_code == 200
    assert (
        await client.post(path, json={**reason, "revision": 1, "cells": cells})
    ).status_code == 409


@pytest.mark.parametrize(
    "value,code",
    [
        ("NaN", "NUMERIC_UNRESOLVED"),
        ("-1", "NEGATIVE_SCORE"),
        ("200", "PERCENT_RANGE"),
        ("75", "PERCENT_CONFLICT"),
    ],
)
def test_percent_validation(value, code):
    rows = [
        ["Name", "Score A", "Score B", "Total", "Maximum", "Percent"],
        ["Fiction A", "2", "3", "5", "10", value],
    ]
    assert code in {i["code"] for i in validate_table(rows)}


def test_no_inferred_formula_and_column_totals():
    assert validate_table([["Name", "Opaque"], ["A", "not a number"]]) == []
    rows = [
        ["Name", "Score A", "Score B", "Total"],
        ["A", "2", "3", "5"],
        ["B", "1", "2", "3"],
        ["Total", "9", "5", "14"],
    ]
    assert "COLUMN_TOTAL_CONFLICT" in {i["code"] for i in validate_table(rows)}


async def test_review_cannot_bypass_document_task(client, adaptive):
    case, _ = await upload(client, "UNKNOWN")
    task = case["review_tasks"][0]
    assert (
        await client.post(
            f"/api/v1/review-tasks/{task['id']}/decision",
            json={"action": "resolve", "reason": "Attempt bypass"},
        )
    ).status_code == 409
    filtered = await client.get("/api/v1/review-tasks?task_type=DOCUMENT_REVIEW")
    assert filtered.status_code == 200 and filtered.json()["total"] == 1
    doc = case["documents"][0]
    result = await client.post(
        f"/api/v1/documents/{doc['id']}/review",
        json={"revision": 0, "reason": "Read original; no structured extraction expected"},
    )
    assert result.status_code == 200 and result.json()["documents"][0]["reviewed"]


async def test_preview_and_summary(client, adaptive):
    case, _ = await upload(client, "PRINTED_TABLE")
    doc = case["documents"][0]
    response = await client.get(f"/api/v1/documents/{doc['id']}/preview")
    assert response.status_code == 200 and response.content.startswith(b"\x89PNG")
    assert (await client.get(f"/api/v1/documents/{doc['id']}/preview?page=2")).status_code == 422
    summary = (await client.get("/api/v1/documents/summary")).json()
    assert summary["document_types"]["PRINTED_TABLE"] == 1


async def test_raster_structure_and_confidence():
    from app.adaptive.layout import CellExtractor

    content, truth = fixture("PRINTED_TABLE", 2)
    evidence = LayoutDetector().native(content, "application/pdf")
    table = evidence.tables[0]
    assert (table.rows, table.columns) == (truth["tables"][0]["row_count"], 6)
    box = table.cells[0].bbox
    CellExtractor().attach(
        evidence.tables,
        [
            {
                "page": 1,
                "width": 600,
                "height": 800,
                "lines": [
                    {
                        "text": "Name",
                        "coordinates": [box[0] + 5, box[1] + 5, 20, 10],
                        "confidence": 0.42,
                    }
                ],
            }
        ],
        evidence.pages,
    )
    assert table.cells[0].value == "Name"
    assert table.cells[0].confidence == 0.42 and table.cells[0].uncertain


async def test_handwriting_provider_preserves_uncertainty():
    from app.adaptive.handwriting import HandwrittenTableExtractor

    class Provider:
        async def recognize(self, image):
            return "7", 0.61, "test-handwriting"

    content, _ = fixture("HANDWRITTEN_TABLE")
    evidence = LayoutDetector().native(content, "application/pdf")
    result = await HandwrittenTableExtractor(Provider()).extract(evidence, {})
    body = [c for c in result.tables[0].cells if c.row > 0]
    assert all(c.value == "7" and c.confidence == 0.61 and c.uncertain for c in body)


async def test_search_filters_and_table_export(client, adaptive):
    import json
    from io import BytesIO

    from openpyxl import load_workbook

    case, _ = await upload(client, "PRINTED_TABLE")
    response = await client.get(
        "/api/v1/cases", params={"q": case["public_id"], "status": "REVIEW_REQUIRED"}
    )
    assert response.json()["total"] == 1
    assert response.json()["items"][0]["document_types"] == ["PRINTED_TABLE"]
    assert response.json()["items"][0]["security_states"] == ["SAFE"]
    assert (await client.get("/api/v1/cases", params={"q": "%"})).json()["total"] == 0
    doc = case["documents"][0]
    result = await client.post(
        f"/api/v1/documents/{doc['id']}/review",
        json={"revision": 0, "reason": "Synthetic export verification", "cells": []},
    )
    assert result.status_code == 200
    assert (await client.post(f"/api/v1/review/{case['id']}/approve")).status_code == 200
    exported = await client.post(f"/api/v1/exports/{case['id']}/json")
    assert json.loads(exported.content)["documents"][0]["tables"][0]["cells"]
    exported = await client.post(f"/api/v1/exports/{case['id']}/xlsx")
    workbook = load_workbook(BytesIO(exported.content))
    assert workbook["Table 1"]["A1"].value == "Name"
    assert workbook["Table 1"]["A2"].value == "Person A"
