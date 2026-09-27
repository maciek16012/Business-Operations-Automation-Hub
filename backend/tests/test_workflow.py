import hashlib
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

SAMPLES = Path(__file__).resolve().parents[2] / "sample_data"


async def create(client, data=None):
    response = await client.post("/api/v1/cases", json=data or {})
    assert response.status_code == 201, response.text
    return response.json()


async def upload(client, case_id, content, filename="document.txt"):
    response = await client.post(
        f"/api/v1/uploads/{case_id}", files=[("files", (filename, content, "text/plain"))]
    )
    assert response.status_code == 200, response.text
    return response.json()


async def test_full_review_and_exports(client):
    case = await create(client)
    identifier = case["id"]
    assert case["status"] == "RECEIVED"
    assert case["public_id"].startswith("CASE-")
    content = (SAMPLES / "review_required.txt").read_bytes()
    case = await upload(client, identifier, content)
    assert case["status"] == "REVIEW_REQUIRED"
    assert case["estimated_value"] == "-500.00"
    assert case["attachments"][0]["sha256"] == hashlib.sha256(content).hexdigest()
    attachment_id = case["attachments"][0]["id"]
    downloaded = await client.get(f"/api/v1/uploads/{identifier}/{attachment_id}")
    assert downloaded.content == content
    blocked = await client.post(f"/api/v1/review/{identifier}/approve")
    assert blocked.status_code == 409
    assert blocked.json()["error"]["code"] == "APPROVAL_BLOCKED"
    assert (await client.post(f"/api/v1/exports/{identifier}/json")).status_code == 409
    original_fields = case["extracted_fields"]
    response = await client.patch(
        f"/api/v1/review/{identifier}",
        json={"tax_id": "526-025-02-74", "estimated_value": "12 500,00 PLN"},
    )
    assert response.status_code == 200, response.text
    case = response.json()
    assert case["status"] == "READY"
    assert case["estimated_value"] == "12500.00"
    assert case["extracted_fields"] == original_fields
    assert all(i["resolved"] for i in case["validation_issues"])
    approved = await client.post(f"/api/v1/review/{identifier}/approve")
    assert approved.json()["status"] == "APPROVED"
    assert (
        await client.patch(f"/api/v1/review/{identifier}", json={"estimated_value": "-1"})
    ).status_code == 409
    assert (
        await client.post(f"/api/v1/review/{identifier}/reject", json={"reason": "late"})
    ).status_code == 409
    exported = await client.post(f"/api/v1/exports/{identifier}/json")
    assert exported.status_code == 201, exported.text
    assert exported.json()["data"]["estimated_value"] == "12500.00"
    assert exported.json()["schema_version"] == 1
    assert "normalization_errors" not in exported.json()
    saved = await client.get(exported.headers["location"])
    assert saved.content == exported.content
    xlsx = await client.post(f"/api/v1/exports/{identifier}/xlsx")
    assert xlsx.status_code == 201
    book = load_workbook(BytesIO(xlsx.content))
    assert book.sheetnames == ["Summary", "Attachments", "Validation", "Audit"]
    summary = dict(book["Summary"].values)
    assert summary["Case ID"] == case["public_id"]
    assert summary["Estimated Value"] == 12500
    assert summary["Status"] == "EXPORTED"
    assert book["Summary"].freeze_panes == "A2"
    assert book["Attachments"].max_row == 2
    final = (await client.get(f"/api/v1/cases/{identifier}")).json()
    assert final["status"] == "EXPORTED" and len(final["exports"]) == 2
    events = [a["event_type"] for a in final["audit_events"]]
    for event in [
        "case_created",
        "attachment_uploaded",
        "extraction_performed",
        "validation_performed",
        "field_manually_corrected",
        "validation_issue_resolved",
        "case_approved",
        "export_generated",
        "status_changed",
    ]:
        assert event in events
    assert events[0] == "case_created" and events[-1] == "export_generated"
    assert (await client.get(f"/api/v1/cases/{identifier}/audit")).json() == final["audit_events"]
    listing = (await client.get("/api/v1/cases")).json()
    assert listing["total"] == 1 and listing["items"][0]["id"] == identifier


async def test_duplicate_renamed_and_scope(client):
    content = b"Customer: Example\nTitle: Quote"
    identifier = (await create(client))["id"]
    case = await upload(client, identifier, content)
    case = await upload(client, identifier, content, "renamed.txt")
    assert len(case["attachments"]) == 1
    assert len(case["extracted_fields"]) == 2
    assert any(i["code"] == "DUPLICATE_ATTACHMENT" for i in case["validation_issues"])
    assert any(e["event_type"] == "duplicate_detected" for e in case["audit_events"])
    another = (await create(client))["id"]
    assert len((await upload(client, another, content))["attachments"]) == 1


async def test_multi_upload_conflict_requires_review(client):
    identifier = (await create(client))["id"]
    response = await client.post(
        f"/api/v1/uploads/{identifier}",
        files=[
            ("files", ("one.txt", b"Customer: Example\nTitle: A", "text/plain")),
            ("files", ("two.txt", b"Title: B", "text/plain")),
            ("files", ("copy.txt", b"Title: B", "text/plain")),
        ],
    )
    assert response.status_code == 200
    case = response.json()
    assert case["status"] == "REVIEW_REQUIRED"
    assert case["request_title"] == "A"
    assert len(case["attachments"]) == 2
    case = (await client.patch(f"/api/v1/review/{identifier}", json={"request_title": "B"})).json()
    assert case["status"] == "READY"


@pytest.mark.parametrize(
    "field,value",
    [
        ("estimated_value", "NaN"),
        ("estimated_value", "1.0001"),
        ("requested_deadline", "tomorrow"),
        ("requested_deadline", "2020-01-01"),
        ("currency", "XYZ"),
        ("customer_email", "broken"),
        ("tax_id", "1234567890"),
    ],
)
async def test_invalid_data_cannot_bypass_api(client, field, value):
    case = await create(
        client, {"customer_name": "Example", "request_title": "Quote", field: value}
    )
    response = await client.post(f"/api/v1/review/{case['id']}/approve")
    assert response.status_code == 409
    detail = (await client.get(f"/api/v1/cases/{case['id']}")).json()
    assert detail["status"] == "REVIEW_REQUIRED"
    assert any(not i["resolved"] and i["severity"] == "ERROR" for i in detail["validation_issues"])


async def test_reject_and_failed_processing(client):
    identifier = (await create(client))["id"]
    assert (
        await client.post(f"/api/v1/review/{identifier}/reject", json={"reason": " "})
    ).status_code == 422
    rejected = await client.post(
        f"/api/v1/review/{identifier}/reject", json={"reason": "Cancelled by customer"}
    )
    assert rejected.json()["status"] == "FAILED"
    assert rejected.json()["review_decisions"][0]["comment"] == "Cancelled by customer"
    assert (await client.post(f"/api/v1/review/{identifier}/approve")).status_code == 409
    identifier = (await create(client))["id"]
    failed = await upload(client, identifier, b"\xff\xfe")
    assert failed["status"] == "FAILED"
    assert len(failed["attachments"]) == 1
    assert any(e["event_type"] == "processing_failure" for e in failed["audit_events"])


async def test_upload_errors_and_structured_requests(client):
    identifier = (await create(client))["id"]
    assert (
        await client.post(
            f"/api/v1/uploads/{identifier}", files={"files": ("a.pdf", b"pdf", "application/pdf")}
        )
    ).status_code == 415
    assert (
        await client.post(
            f"/api/v1/uploads/{identifier}", files={"files": ("a.txt", b"", "text/plain")}
        )
    ).status_code == 413
    assert (
        await client.post(
            f"/api/v1/uploads/{identifier}",
            files={"files": ("a.txt", b"x" * (5 * 1024 * 1024 + 1), "text/plain")},
        )
    ).status_code == 413
    assert (await client.post("/api/v1/cases", json={"status": "APPROVED"})).status_code == 422
    assert (await client.post("/api/v1/cases", json={"source": "email"})).status_code == 422
    assert (await client.get("/api/v1/cases/not-a-uuid")).status_code == 422
    assert (
        await client.get("/api/v1/cases/00000000-0000-0000-0000-000000000000")
    ).status_code == 404
    case = await upload(client, identifier, b"Customer: Example\nTitle: Quote", "../../safe.txt")
    assert case["attachments"][0]["original_filename"] == "safe.txt"
    assert case["status"] == "READY"


async def test_formula_injection_is_plain_text(client):
    case = await create(client, {"customer_name": "=1+1", "request_title": "+formula"})
    identifier = case["id"]
    assert (await client.post(f"/api/v1/review/{identifier}/approve")).status_code == 200
    exported = await client.post(f"/api/v1/exports/{identifier}/xlsx")
    book = load_workbook(BytesIO(exported.content))
    assert book["Summary"]["B3"].value == "=1+1"
    assert book["Summary"]["B3"].data_type == "s"


async def test_valid_fixture(client):
    identifier = (await create(client))["id"]
    case = await upload(client, identifier, (SAMPLES / "valid_inquiry.txt").read_bytes())
    assert case["status"] == "READY"
    assert case["currency"] == "PLN" and case["estimated_value"] == "12500.00"
    assert (await client.post(f"/api/v1/review/{identifier}/approve")).status_code == 200


async def test_export_storage_failure_is_audited_and_rolls_back(client, monkeypatch):
    from app.storage.base import LocalFilesystemStorage

    case = await create(client, {"customer_name": "Example", "request_title": "Quote"})
    identifier = case["id"]
    assert (await client.post(f"/api/v1/review/{identifier}/approve")).status_code == 200

    def unavailable(self, content):
        raise OSError("Simulated disk failure")

    monkeypatch.setattr(LocalFilesystemStorage, "put", unavailable)
    response = await client.post(f"/api/v1/exports/{identifier}/json")
    assert response.status_code == 503
    case = (await client.get(f"/api/v1/cases/{identifier}")).json()
    assert case["status"] == "APPROVED"
    assert case["exports"][0]["metadata_json"] == {"success": False}
    assert not any(e["event_type"] == "export_generated" for e in case["audit_events"])
    assert any(e["event_type"] == "processing_failure" for e in case["audit_events"])


async def test_upload_storage_failure_is_audited(client, monkeypatch):
    from app.storage.base import LocalFilesystemStorage

    identifier = (await create(client))["id"]

    def unavailable(self, content):
        raise OSError("Simulated disk failure")

    monkeypatch.setattr(LocalFilesystemStorage, "put", unavailable)
    case = await upload(client, identifier, b"Customer: Example\nTitle: Quote")
    assert case["status"] == "FAILED"
    assert any(e["event_type"] == "processing_failure" for e in case["audit_events"])


async def test_currency_conflict_and_malformed_correction(client):
    identifier = (await create(client))["id"]
    case = await upload(
        client,
        identifier,
        b"Customer: Example\nTitle: Quote\nCurrency: EUR\nEstimated value: 100 PLN",
    )
    assert case["status"] == "REVIEW_REQUIRED"
    corrected = await client.patch(
        f"/api/v1/review/{identifier}", json={"currency": "PLN", "estimated_value": "100 PLN"}
    )
    assert corrected.json()["status"] == "READY"
    malformed = await client.patch(f"/api/v1/review/{identifier}", json={"estimated_value": "bad"})
    assert malformed.json()["status"] == "REVIEW_REQUIRED"
    assert malformed.json()["estimated_value"] is None
    assert (await client.post(f"/api/v1/review/{identifier}/approve")).status_code == 409
    cleared = await client.patch(f"/api/v1/review/{identifier}", json={"estimated_value": None})
    assert cleared.json()["status"] == "READY"


async def test_revalidation_preserves_unchanged_issues(client):
    identifier = (await create(client))["id"]
    first = (await client.post(f"/api/v1/review/{identifier}/validate")).json()
    second = (await client.post(f"/api/v1/review/{identifier}/validate")).json()
    assert first["validation_issues"] == second["validation_issues"]
    assert (await client.post(f"/api/v1/exports/{identifier}/pdf")).status_code == 422


async def test_postgres_concurrent_duplicate_and_approval(client):
    import asyncio
    import os

    if not os.getenv("TEST_DATABASE_URL", "").startswith("postgresql"):
        pytest.skip("Row locking requires PostgreSQL; run with TEST_DATABASE_URL")
    identifier = (await create(client))["id"]
    content = b"Customer: Example\nTitle: Concurrent uploads"
    results = await asyncio.gather(
        upload(client, identifier, content), upload(client, identifier, content, "renamed.txt")
    )
    assert all(r["status"] == "READY" for r in results)
    case = (await client.get(f"/api/v1/cases/{identifier}")).json()
    assert len(case["attachments"]) == 1
    assert len([e for e in case["audit_events"] if e["event_type"] == "duplicate_detected"]) == 1
    approvals = await asyncio.gather(
        client.post(f"/api/v1/review/{identifier}/approve"),
        client.post(f"/api/v1/review/{identifier}/approve"),
    )
    assert sorted(r.status_code for r in approvals) == [200, 409]


async def test_control_characters_rejected_without_database_failure(client):
    assert (
        await client.post("/api/v1/cases", json={"request_title": "bad\x00title"})
    ).status_code == 422
    identifier = (await create(client))["id"]
    failed = await upload(client, identifier, b"Customer: Example\nTitle: Bad\x00title")
    assert failed["status"] == "FAILED"
    assert any(e["event_type"] == "processing_failure" for e in failed["audit_events"])


async def test_inferred_manual_currency_is_audited(client):
    identifier = (await create(client, {"customer_name": "Example", "request_title": "Quote"}))[
        "id"
    ]
    response = await client.patch(
        f"/api/v1/review/{identifier}", json={"estimated_value": "100 PLN"}
    )
    assert response.status_code == 200
    case = response.json()
    assert case["currency"] == "PLN"
    changes = [
        e["details"] for e in case["audit_events"] if e["event_type"] == "field_manually_corrected"
    ]
    assert any(c["field"] == "currency" and c["old"] is None and c["new"] == "PLN" for c in changes)
