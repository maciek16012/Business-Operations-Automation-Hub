import asyncio
import base64
import copy
import hashlib
import os
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.storage.base import LocalFilesystemStorage


def email(message_id="<unit-email-001@example.test>"):
    return {
        "external_message_id": message_id,
        "sender": {"address": "jan@example.com", "name": "Jan Example"},
        "recipients": [{"address": "office@example.com", "name": "Office"}],
        "cc": [{"address": "cc@example.com"}],
        "reply_to": [{"address": "reply@example.com"}],
        "subject": "Synthetic inquiry",
        "received_at": "2026-09-27T10:00:00+02:00",
        "sent_at": "2026-09-27T07:59:00Z",
        "text_body": "Please prepare a quote.\nSynthetic data.",
        "html_body": '<script>alert("untrusted")</script><p>Quote</p>',
        "attachments": [],
    }


def attachment(
    content=b"NIP: 5260250274\nEstimated value: 12 500,00 PLN",
    filename="quote.txt",
    mime="text/plain",
):
    return {
        "filename": filename,
        "mime_type": mime,
        "content_base64": base64.b64encode(content).decode(),
    }


async def deliver(client, payload):
    response = await client.post("/api/v1/inbound/email", json=payload)
    assert response.status_code in {200, 201}, response.text
    return response


async def detail(client, result):
    return (await client.get(f"/api/v1/cases/{result['case_id']}")).json()


async def test_metadata_zero_attachments_and_source(client):
    payload = email()
    response = await deliver(client, payload)
    assert response.status_code == 201
    result = response.json()
    assert result["result"] == "created" and result["status"] == "READY"
    assert result["processing_status"] == "processed" and result["attachment_results"] == []
    case = await detail(client, result)
    assert case["source"] == "email" and case["attachments"] == []
    assert case["customer_name"] == "Jan Example" and case["customer_email"] == "jan@example.com"
    assert case["request_title"] == payload["subject"]
    source = case["inbound_message"]
    for key in ("subject", "text_body", "html_body", "external_message_id"):
        assert source[key] == payload[key]
    assert source["sender_address"] == payload["sender"]["address"]
    assert source["received_at"].startswith("2026-09-27T08:00:00")
    assert source["recipients"] == [{"address": "office@example.com", "name": "Office"}]
    assert source["cc"][0]["address"] == "cc@example.com"
    assert source["reply_to"][0]["address"] == "reply@example.com"
    events = {e["event_type"] for e in case["audit_events"]}
    assert {"EMAIL_RECEIVED", "EMAIL_CASE_CREATED", "validation_performed"} <= events


async def test_single_attachment_replay_preserves_exactly_one_case(client):
    payload = email()
    payload["attachments"] = [attachment()]
    original = (await deliver(client, payload)).json()
    case_before = await detail(client, original)
    replay = copy.deepcopy(payload)
    replay["received_at"] = "2026-10-01T12:00:00Z"
    replay["attachments"][0]["filename"] = "renamed.txt"
    for _ in range(3):
        response = await deliver(client, replay)
        assert response.status_code == 200
        result = response.json()
        assert result["result"] == "duplicate"
        assert (
            result["case_id"] == original["case_id"]
            and result["message_id"] == original["message_id"]
        )
    case = await detail(client, original)
    assert case["status"] == "READY" and case["estimated_value"] == "12500.00"
    assert case["attachments"] == case_before["attachments"]
    assert case["extracted_fields"] == case_before["extracted_fields"]
    assert len(case["attachments"]) == 1
    assert (
        case["attachments"][0]["sha256"]
        == hashlib.sha256(base64.b64decode(payload["attachments"][0]["content_base64"])).hexdigest()
    )
    assert (await client.get("/api/v1/cases")).json()["total"] == 1
    events = [e["event_type"] for e in case["audit_events"]]
    assert events.count("EMAIL_DUPLICATE_IGNORED") == 3
    assert events.count("extraction_performed") == 1
    assert events.count("EMAIL_ATTACHMENT_STORED") == 1


async def test_multiple_attachments_content_dedupe_and_safe_filenames(client):
    payload = email()
    payload["attachments"] = [
        attachment(filename="../../quote.txt"),
        attachment(filename="renamed.txt"),
        attachment(b"Description: Extra specification", "extra.txt"),
    ]
    result = (await deliver(client, payload)).json()
    assert [a["result"] for a in result["attachment_results"]] == ["stored", "duplicate", "stored"]
    case = await detail(client, result)
    assert len(case["attachments"]) == 2
    assert case["attachments"][0]["original_filename"] == "quote.txt"
    assert result["attachment_results"][0]["original_filename"] == "../../quote.txt"
    assert all(".." not in a["storage_key"] for a in case["attachments"])
    assert any(e["event_type"] == "duplicate_detected" for e in case["audit_events"])


@pytest.mark.parametrize("message_id", [None, "", "not a trustworthy message id"])
async def test_fallback_ignores_receipt_time_filename_and_attachment_order(client, message_id):
    payload = email(message_id)
    payload["attachments"] = [attachment(), attachment(b"Description: More", "extra.txt")]
    first = (await deliver(client, payload)).json()
    assert first["identity_method"] == "fingerprint_v1"
    payload["received_at"] = "2026-12-01T12:00:00Z"
    payload["attachments"].reverse()
    payload["attachments"][0]["filename"] = "new-name.txt"
    replay = (await deliver(client, payload)).json()
    assert replay["result"] == "duplicate" and replay["case_id"] == first["case_id"]
    payload["text_body"] += "A different body"
    changed = (await deliver(client, payload)).json()
    assert changed["result"] == "created" and changed["case_id"] != first["case_id"]


async def test_rfc_identity_normalization_and_distinct_messages(client):
    first = (await deliver(client, email("<Id-123@EXAMPLE.TEST>"))).json()
    second = (await deliver(client, email("  Id-123@example.test  "))).json()
    assert second["case_id"] == first["case_id"] and second["result"] == "duplicate"
    third = (await deliver(client, email("<Id-124@example.test>"))).json()
    assert third["case_id"] != first["case_id"]


async def test_incomplete_email_review_approval_and_exports(client):
    payload = email()
    payload["sender"]["name"] = None
    payload["subject"] = ""
    result = (await deliver(client, payload)).json()
    identifier = result["case_id"]
    assert result["status"] == "REVIEW_REQUIRED"
    assert (await client.post(f"/api/v1/review/{identifier}/approve")).status_code == 409
    corrected = await client.patch(
        f"/api/v1/review/{identifier}",
        json={"customer_name": "Verified Example", "request_title": "Reviewed request"},
    )
    assert corrected.json()["status"] == "READY"
    assert (await client.post(f"/api/v1/review/{identifier}/approve")).json()[
        "status"
    ] == "APPROVED"
    export = await client.post(f"/api/v1/exports/{identifier}/json")
    assert (
        export.status_code == 201 and export.json()["data"]["customer_name"] == "Verified Example"
    )
    xlsx = await client.post(f"/api/v1/exports/{identifier}/xlsx")
    assert load_workbook(BytesIO(xlsx.content)).sheetnames == [
        "Summary",
        "Attachments",
        "Validation",
        "Audit",
    ]
    replay = (await deliver(client, payload)).json()
    assert replay["result"] == "duplicate" and replay["status"] == "EXPORTED"
    assert (await client.get("/api/v1/cases")).json()["total"] == 1


async def test_unsupported_attachment_is_preserved_and_requires_explicit_review(client):
    payload = email()
    original = b"%PDF-1.4 synthetic document"
    payload["attachments"] = [attachment(original, "offer.pdf", "application/pdf"), attachment()]
    result = (await deliver(client, payload)).json()
    assert result["status"] == "REVIEW_REQUIRED"
    assert result["attachment_results"][0]["result"] == "review_required"
    case = await detail(client, result)
    assert len(case["attachments"]) == 2 and case["estimated_value"] == "12500.00"
    artifact = case["attachments"][0]
    response = await client.get(f"/api/v1/uploads/{case['id']}/{artifact['id']}")
    assert response.content == original
    for action in ("validate", "approve"):
        response = await client.post(f"/api/v1/review/{case['id']}/{action}")
        assert response.status_code == (200 if action == "validate" else 409)
    url = f"/api/v1/inbound/messages/{result['message_id']}/attachments/{artifact['id']}/review"
    assert (await client.post(url, json={"reason": " "})).status_code == 422
    reviewed = await client.post(
        url, json={"reason": "Inspected original and confirmed quoted data"}
    )
    assert reviewed.status_code == 200 and reviewed.json()["status"] == "READY"
    assert any(
        e["event_type"] == "EMAIL_ATTACHMENT_REVIEWED" for e in reviewed.json()["audit_events"]
    )
    assert (await client.post(f"/api/v1/review/{case['id']}/approve")).status_code == 200
    assert (await client.post(url, json={"reason": "Repeat"})).status_code == 409


async def test_failed_extraction_does_not_drop_later_attachments(client):
    payload = email()
    payload["attachments"] = [attachment(b"\xff", "broken.txt"), attachment()]
    result = (await deliver(client, payload)).json()
    case = await detail(client, result)
    assert case["status"] == "REVIEW_REQUIRED" and len(case["attachments"]) == 2
    assert case["estimated_value"] == "12500.00"
    assert any(i["code"] == "EMAIL_EXTRACTION_FAILED" for i in case["validation_issues"])
    assert any(e["event_type"] == "EMAIL_INGESTION_FAILED" for e in case["audit_events"])


async def test_storage_failure_is_persisted_once_and_audited(client, monkeypatch):
    def fail(self, content):
        raise OSError("test storage failure")

    monkeypatch.setattr(LocalFilesystemStorage, "put", fail)
    payload = email()
    payload["attachments"] = [attachment()]
    first = (await deliver(client, payload)).json()
    assert first["processing_status"] == "failed" and first["status"] == "FAILED"
    second = (await deliver(client, payload)).json()
    assert second["result"] == "duplicate" and second["case_id"] == first["case_id"]
    case = await detail(client, first)
    assert any(e["event_type"] == "EMAIL_INGESTION_FAILED" for e in case["audit_events"])
    assert (await client.get("/api/v1/cases")).json()["total"] == 1


@pytest.mark.parametrize(
    "change",
    [
        {"sender": {"address": "not-an-email"}},
        {"received_at": "yesterday"},
        {"received_at": "2026-09-27T10:00:00"},
        {"arbitrary_header": "not allowed"},
        {"subject": "bad\x00subject"},
        {"attachments": [attachment()] * 11},
        {"attachments": [{"filename": "bad.txt", "content_base64": "%%%"}]},
        {"attachments": [attachment(b"")]},
    ],
)
async def test_invalid_payload_creates_no_case(client, change):
    payload = email()
    payload.update(change)
    response = await client.post("/api/v1/inbound/email", json=payload)
    assert response.status_code == 422
    assert (await client.get("/api/v1/cases")).json()["total"] == 0


async def test_attachment_size_limit_is_shared(client, monkeypatch):
    from app.core.config import settings

    monkeypatch.setattr(settings, "max_upload_bytes", 4)
    payload = email()
    payload["attachments"] = [attachment(b"12345")]
    assert (await client.post("/api/v1/inbound/email", json=payload)).status_code == 422
    assert (await client.get("/api/v1/cases")).json()["total"] == 0


async def test_concurrent_deliveries_create_one_case_postgres(client):
    if not os.getenv("TEST_DATABASE_URL", "").startswith("postgresql"):
        pytest.skip("Concurrent unique-key arbitration requires PostgreSQL")
    payload = email()
    payload["attachments"] = [attachment()]
    responses = await asyncio.gather(*(deliver(client, payload) for _ in range(6)))
    assert sorted(r.status_code for r in responses) == [200, 200, 200, 200, 200, 201]
    assert len({r.json()["case_id"] for r in responses}) == 1
    case = await detail(client, responses[0].json())
    assert len(case["attachments"]) == 1
    assert sum(e["event_type"] == "extraction_performed" for e in case["audit_events"]) == 1
    assert sum(e["event_type"] == "EMAIL_DUPLICATE_IGNORED" for e in case["audit_events"]) == 5
    assert (await client.get("/api/v1/cases")).json()["total"] == 1
